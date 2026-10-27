from io import BytesIO
from pathlib import Path

import streamlit as st
from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {"pdf", "docx", "txt"}


def extract_text(uploaded_file) -> str:
    """Extract readable text from an uploaded document."""
    extension = Path(uploaded_file.name).suffix.lower().lstrip(".")

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {extension}")

    file_bytes = uploaded_file.getvalue()

    if not file_bytes:
        raise ValueError("The uploaded file is empty.")

    if extension == "pdf":
        reader = PdfReader(BytesIO(file_bytes))

        if reader.is_encrypted:
            raise ValueError(
                "Encrypted PDFs are not supported. "
                "Please upload an accessible PDF."
            )

        pages = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"[Page {page_number}]\n{text}")

        extracted = "\n\n".join(pages)

    elif extension == "docx":
        document = Document(BytesIO(file_bytes))
        paragraphs = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        table_rows = []
        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip().replace("\n", " ")
                    for cell in row.cells
                ]
                table_rows.append(" | ".join(cells))

        extracted = "\n".join(paragraphs)
        if table_rows:
            extracted += "\n\nTable content:\n" + "\n".join(table_rows)

    else:
        try:
            extracted = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError:
            extracted = file_bytes.decode("utf-8", errors="replace")

    if not extracted.strip():
        raise ValueError(
            "No readable text was found. "
            "Scanned PDFs may require OCR before processing."
        )

    return extracted.strip()